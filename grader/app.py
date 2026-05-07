from flask import Flask, render_template, request, send_file, jsonify, Response, stream_with_context
import pandas as pd
from validate_infrastructure import InfrastructureGrader
import os
import json
import time
from datetime import datetime
import queue
import threading
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment

app = Flask(__name__)

# Paths
BASE_DIR = os.path.dirname(__file__)
EXPORTS_DIR = os.path.join(BASE_DIR, "exports")
DB_FILE = os.path.join(BASE_DIR, "database.json")

if not os.path.exists(EXPORTS_DIR): os.makedirs(EXPORTS_DIR)
if not os.path.exists(DB_FILE):
    with open(DB_FILE, "w") as f: json.dump([], f)

def get_db():
    try:
        with open(DB_FILE, "r") as f:
            content = f.read()
            return json.loads(content) if content else []
    except:
        return []

def save_to_db(entry):
    db = get_db()
    db.append(entry)
    with open(DB_FILE, "w") as f:
        json.dump(db, f, indent=4)

@app.route("/")
def index():
    return """
    <!DOCTYPE html>
    <html lang="en">
    <head>
        <meta charset="UTF-8">
        <meta name="viewport" content="width=device-width, initial-scale=1.0">
        <title>NeonStage Infrastructure Grader</title>
        <link href="https://fonts.googleapis.com/css2?family=Outfit:wght@300;400;600;700&display=swap" rel="stylesheet">
        <style>
            :root {
                --primary: #818cf8; --bg: #020617; --card: #1e293b; --text: #f8fafc;
                --accent: #38bdf8; --success: #4ade80; --fail: #fb7185;
            }
            body { font-family: 'Outfit', sans-serif; background: var(--bg); color: var(--text); margin: 0; padding: 20px; line-height: 1.5; }
            .container { max-width: 1400px; margin: 0 auto; }
            .grid { display: grid; grid-template-columns: 400px 1fr; gap: 20px; }
            .card { background: var(--card); padding: 25px; border-radius: 16px; box-shadow: 0 10px 15px -3px rgba(0,0,0,0.1); margin-bottom: 20px; border: 1px solid #334155; }
            h1, h2, h3 { color: var(--primary); margin-top: 0; }
            .form-group { margin-bottom: 15px; }
            input { width: 100%; padding: 12px; border-radius: 8px; border: 1px solid #334155; background: #0f172a; color: white; box-sizing: border-box; font-size: 1rem; }
            button { width: 100%; padding: 14px; background: var(--primary); color: white; border: none; border-radius: 8px; font-weight: 700; cursor: pointer; transition: 0.3s; font-size: 1rem; }
            button:hover { background: #6366f1; transform: translateY(-2px); }
            button.secondary { background: #475569; }
            button.delete { background: var(--fail); width: auto; padding: 5px 10px; font-size: 0.75rem; }
            
            table { width: 100%; border-collapse: collapse; margin-top: 10px; }
            th, td { padding: 12px; text-align: left; border-bottom: 1px solid #334155; }
            th { color: #94a3b8; text-transform: uppercase; font-size: 0.75rem; letter-spacing: 0.1em; background: #0f172a; position: sticky; top: 0; }
            
            .status-pass { color: var(--success); font-weight: 700; }
            .status-fail { color: var(--fail); font-weight: 700; }
            .score-pill { background: var(--primary); padding: 4px 12px; border-radius: 20px; font-weight: 700; }
            
            .loader-box { display: none; margin-top: 20px; text-align: center; }
            .progress-text { font-weight: 600; color: var(--accent); margin-top: 10px; font-size: 0.9rem; min-height: 1.5em; }
            
            .spinner {
                width: 40px; height: 40px; border: 4px solid rgba(255,255,255,0.1); border-left-color: var(--primary);
                border-radius: 50%; animation: spin 1s linear infinite; margin: 0 auto;
            }
            @keyframes spin { to { transform: rotate(360deg); } }

            .scroll-box { max-height: 700px; overflow-y: auto; border-radius: 8px; background: #0f172a; border: 1px solid #334155; }
            .badge { background: #334155; padding: 2px 8px; border-radius: 4px; font-size: 0.7rem; color: #94a3b8; }
        </style>
    </head>
    <body>
        <div class="container">
            <h1 style="display:flex; align-items:center; gap:15px; margin-bottom:40px;">
                <span style="font-size:3rem">🎯</span> 
                <span>NeonStage <span style="color:var(--accent)">Infrastructure Grader</span> <span style="font-size:1rem; opacity:0.5">v2.1</span></span>
            </h1>

            <div class="grid">
                <div class="sidebar">
                    <div class="card">
                        <h2>New Evaluation</h2>
                        <form id="gradeForm">
                            <div class="form-group"><input type="text" id="student_name" placeholder="Student Full Name" required></div>
                            <div class="form-group"><input type="text" id="access_key" placeholder="AWS Access Key" required></div>
                            <div class="form-group"><input type="password" id="secret_key" placeholder="AWS Secret Key" required></div>
                            <div class="input-group">
                                <label for="session_token">Session Token</label>
                                <input type="text" id="session_token" placeholder="Optional for LabRole">
                            </div>
                            
                            <button type="submit" id="submitBtn">Run 280+ Point Scan</button>
                        </form>
                        <div class="loader-box" id="loaderBox">
                            <div class="spinner"></div>
                            <div class="progress-text" id="progressText">Initializing...</div>
                        </div>
                    </div>

                    <div class="card">
                        <div style="display:flex; justify-content:space-between; align-items:center; margin-bottom:15px;">
                            <h2>Leaderboard</h2>
                            <div style="display:flex; gap:10px;">
                                <button onclick="exportAll()" class="primary" style="width:auto; padding:5px 15px; font-size:0.8rem; background:var(--primary)">Export Master Analytic</button>
                                <button onclick="exportHtml()" class="secondary" style="width:auto; padding:5px 15px; font-size:0.8rem; background:#0ea5e9; color:white; border:none;">Export Static HTML</button>
                                <button onclick="loadLeaderboard()" class="secondary" style="width:auto; padding:5px 15px; font-size:0.8rem">Refresh</button>
                                <button onclick="resetLeaderboard()" class="delete" style="width:auto; padding:5px 15px; font-size:0.8rem">Reset</button>
                            </div>
                        </div>
                        <div class="scroll-box" style="max-height: 400px;">
                            <table id="leaderboardTable">
                                <thead><tr><th>Rank</th><th>Name</th><th>Score</th><th>Action</th></tr></thead>
                                <tbody></tbody>
                            </table>
                        </div>
                    </div>
                </div>

                <div class="main">
                    <!-- Manual Checklist Section (Initially Hidden) -->
                    <div id="manualCheckCard" class="card" style="display:none; border: 2px solid var(--primary); background: rgba(129, 140, 248, 0.05);">
                        <h2 style="color:var(--primary); margin-top:0">Step 2: Manual Web Verification</h2>
                        <p style="font-size:0.9rem; color:#94a3b8">The infrastructure scan is complete. Now, click the ALB link below and verify the frontend functionality.</p>
                        
                        <div style="background:var(--card); padding:15px; border-radius:8px; margin-bottom:10px; border:1px solid var(--primary);">
                            <label style="display:block; font-size:0.8rem; color:#94a3b8; margin-bottom:5px">ALB URL (Application Interface):</label>
                            <a id="albLink" href="#" target="_blank" style="color:var(--accent); font-weight:700; word-break:break-all"></a>
                        </div>
                        
                        <div style="background:var(--card); padding:15px; border-radius:8px; margin-bottom:20px; border:1px solid var(--accent);">
                            <label style="display:block; font-size:0.8rem; color:#94a3b8; margin-bottom:5px">API Gateway URL (Copy to Frontend):</label>
                            <code id="apiLink" style="color:var(--success); font-weight:700; word-break:break-all; background:transparent"></code>
                        </div>

                        <div id="checklistItems">
                            <div style="margin-bottom:10px; display:flex; align-items:center; gap:10px;">
                                <input type="checkbox" id="check_rooms" style="width:20px; height:20px">
                                <label>Room list is visible and correct</label>
                            </div>
                            <div style="margin-bottom:10px; display:flex; align-items:center; gap:10px;">
                                <input type="checkbox" id="check_booking" style="width:20px; height:20px">
                                <label>Booking process (atomic lock) works</label>
                            </div>
                            <div style="margin-bottom:10px; display:flex; align-items:center; gap:10px;">
                                <input type="checkbox" id="check_payment" style="width:20px; height:20px">
                                <label>Payment proof upload (S3) works</label>
                            </div>
                            <div style="margin-bottom:10px; display:flex; align-items:center; gap:10px;">
                                <input type="checkbox" id="check_status" style="width:20px; height:20px">
                                <label>Booking status tracking works</label>
                            </div>
                        </div>

                        <button onclick="finalizeGrade()" class="primary" style="margin-top:20px">Finalize & Save Grade</button>
                    </div>

                    <div id="resultCard" class="card" style="display:none;">
                        <div style="display:flex; justify-content:space-between; align-items:center;">
                            <h2 id="resName" style="margin:0">Scan Results</h2>
                            <button id="exportBtn" style="width:auto; padding: 10px 25px;">Export Detailed Excel</button>
                        </div>
                        <div style="margin: 20px 0; display:flex; gap:20px; align-items:center;">
                            <div class="score-pill" id="resScore" style="font-size:1.5rem">0 / 0</div>
                            <div id="resPercent" style="font-size:1.2rem; color:var(--accent); font-weight:700">0%</div>
                        </div>
                        <div class="scroll-box">
                            <table>
                                <thead>
                                    <tr>
                                        <th>Category</th>
                                        <th>Item</th>
                                        <th>Status</th>
                                        <th>Explanation / Discrepancy</th>
                                    </tr>
                                </thead>
                                <tbody id="resultBody"></tbody>
                            </table>
                        </div>
                    </div>
                </div>
            </div>
        </div>

        <script>
            let currentResults = [];
            let currentStudent = "";
            let currentMetadata = {};
            let fullDb = [];

            async function finalizeGrade() {
                const checklist = {
                    "Rooms List": document.getElementById('check_rooms').checked,
                    "Booking Logic": document.getElementById('check_booking').checked,
                    "Payment Upload": document.getElementById('check_payment').checked,
                    "Status Tracking": document.getElementById('check_status').checked
                };
                
                const manualResults = Object.entries(checklist).map(([item, pass]) => ({
                    Category: "9. Manual Check",
                    Item: item,
                    Status: pass ? "PASS" : "FAIL",
                    Score: pass ? 1 : 0,
                    Expected: "Functional",
                    Actual: pass ? "Functional" : "Not Tested/Failed",
                    Feedback: "Manual verification"
                }));
                
                const finalResults = [...currentResults, ...manualResults];
                const total_score = finalResults.reduce((acc, r) => acc + r.Score, 0);
                const max_score = finalResults.length;
                const percentage = ((total_score / max_score) * 100).toFixed(2);
                
                const entry = {
                    id: Date.now().toString(),
                    name: currentStudent,
                    total_score,
                    max_score,
                    percentage,
                    results: finalResults
                };
                
                await fetch('/api/finalize', {
                    method: 'POST',
                    headers: {'Content-Type': 'application/json'},
                    body: JSON.stringify(entry)
                });
                
                alert("Grade Finalized!");
                document.getElementById('manualCheckCard').style.display = 'none';
                showResult(entry, currentStudent);
                loadLeaderboard();
            }

            async function loadLeaderboard() {
                const res = await fetch('/api/leaderboard');
                const data = await res.json();
                fullDb = data;
                const tbody = document.querySelector('#leaderboardTable tbody');
                tbody.innerHTML = "";
                data.forEach((s) => {
                    const row = document.createElement('tr');
                    row.innerHTML = `
                        <td><span class="badge" style="background:var(--primary)">#${s.rank}</span></td>
                        <td><a href="javascript:void(0)" onclick="viewDetails('${s.id}')" style="color:var(--accent); font-weight:700; text-decoration:none">${s.name}</a></td>
                        <td>${s.percentage}%</td>
                        <td><button class="delete" onclick="deleteEntry('${s.id}')">Delete</button></td>
                    `;
                    tbody.appendChild(row);
                });
            }

            function viewDetails(id) {
                const entry = fullDb.find(e => e.id === id);
                if(entry) {
                    currentResults = entry.results;
                    currentStudent = entry.name;
                    showResult(entry, entry.name);
                    window.scrollTo({top: document.getElementById('resultCard').offsetTop - 20, behavior: 'smooth'});
                }
            }

            async function exportAll() {
                const btn = event.target;
                const originalText = btn.innerText;
                btn.innerText = "Exporting...";
                btn.disabled = true;
                
                try {
                    const res = await fetch('/api/export', { method: 'POST' });
                    if(!res.ok) throw new Error("Export failed");
                    const blob = await res.blob();
                    const url = window.URL.createObjectURL(blob);
                    const a = document.createElement('a');
                    a.href = url;
                    a.download = `Analytical_Master_Report_${new Date().toISOString().slice(0,10)}.xlsx`;
                    document.body.appendChild(a);
                    a.click();
                    a.remove();
                } catch(e) {
                    alert(e.message);
                } finally {
                    btn.innerText = originalText;
                    btn.disabled = false;
                }
            }

            async function exportHtml() {
                const btn = event.target;
                const originalText = btn.innerText;
                btn.innerText = "Generating...";
                btn.disabled = true;
                
                try {
                    const res = await fetch('/api/export-html', { method: 'POST' });
                    if(!res.ok) throw new Error("Export failed");
                    const blob = await res.blob();
                    const url = window.URL.createObjectURL(blob);
                    const a = document.createElement('a');
                    a.href = url;
                    a.download = `Static_Leaderboard_${new Date().toISOString().slice(0,10)}.html`;
                    document.body.appendChild(a);
                    a.click();
                    a.remove();
                } catch(e) {
                    alert(e.message);
                } finally {
                    btn.innerText = originalText;
                    btn.disabled = false;
                }
            }

            async function resetLeaderboard() {
                if(!confirm("DELETE ALL DATA? This cannot be undone.")) return;
                await fetch('/api/reset', { method: 'POST' });
                loadLeaderboard();
            }

            async function deleteEntry(id) {
                if(!confirm("Are you sure?")) return;
                await fetch(`/api/delete/${id}`, { method: 'DELETE' });
                loadLeaderboard();
            }

            document.getElementById('gradeForm').addEventListener('submit', async (e) => {
                e.preventDefault();
                const loaderBox = document.getElementById('loaderBox');
                const progressText = document.getElementById('progressText');
                const submitBtn = document.getElementById('submitBtn');
                const resultCard = document.getElementById('resultCard');
                const manualCheckCard = document.getElementById('manualCheckCard');
                
                loaderBox.style.display = 'block';
                resultCard.style.display = 'none';
                manualCheckCard.style.display = 'none';
                submitBtn.disabled = true;
                progressText.innerText = "Scanning Infrastructure & API...";

                const student_name = document.getElementById('student_name').value;
                const access_key = document.getElementById('access_key').value;
                const secret_key = document.getElementById('secret_key').value;
                const session_token = document.getElementById('session_token').value;

                // Use SSE for progress and final result
                const params = new URLSearchParams({ student_name, access_key, secret_key, session_token });
                const eventSource = new EventSource(`/api/grade/stream?${params.toString()}`);

                eventSource.onmessage = (event) => {
                    const data = JSON.parse(event.data);
                    
                    if (data.type === 'progress') {
                        progressText.innerText = data.message;
                    } else if (data.type === 'complete') {
                        currentResults = data.results;
                        currentStudent = student_name;
                        currentMetadata = data.metadata;
                        
                        // Show Manual Checklist
                        manualCheckCard.style.display = 'block';
                        document.getElementById('albLink').innerText = data.metadata.alb_url;
                        document.getElementById('albLink').href = "http://" + data.metadata.alb_url;
                        document.getElementById('apiLink').innerText = data.metadata.api_url;
                        
                        eventSource.close();
                        loaderBox.style.display = 'none';
                        submitBtn.disabled = false;
                        window.scrollTo({top: manualCheckCard.offsetTop - 20, behavior: 'smooth'});
                    } else if (data.type === 'error') {
                        alert("AWS Error: " + data.message);
                        eventSource.close();
                        loaderBox.style.display = 'none';
                        submitBtn.disabled = false;
                    }
                };

                eventSource.onerror = (err) => {
                    alert("System Error: Connection lost");
                    eventSource.close();
                    loaderBox.style.display = 'none';
                    submitBtn.disabled = false;
                };
            });

            function showResult(data, name) {
                document.getElementById('resultCard').style.display = 'block';
                document.getElementById('resName').innerText = name;
                document.getElementById('resScore').innerText = `${data.total_score} / ${data.max_score}`;
                document.getElementById('resPercent').innerText = `${data.percentage}%`;
                
                const body = document.getElementById('resultBody');
                body.innerHTML = "";
                data.results.forEach(r => {
                    const row = document.createElement('tr');
                    row.innerHTML = `
                        <td><span class="badge">${r.Category}</span></td>
                        <td style="font-size:0.9rem">${r.Item}</td>
                        <td class="${r.Status === 'PASS' ? 'status-pass' : 'status-fail'}">${r.Status}</td>
                        <td style="font-size:0.8rem; color:#94a3b8; font-style:italic">${r.Feedback}</td>
                    `;
                    body.appendChild(row);
                });
            }

            document.getElementById('exportBtn').addEventListener('click', async () => {
                const response = await fetch('/api/export', {
                    method: 'POST',
                    headers: {'Content-Type': 'application/json'},
                    body: JSON.stringify({ results: currentResults, name: currentStudent })
                });
                const blob = await response.blob();
                const url = window.URL.createObjectURL(blob);
                const a = document.createElement('a');
                a.href = url;
                a.download = `MasterReport_${currentStudent}.xlsx`;
                document.body.appendChild(a);
                a.click();
            });

            loadLeaderboard();
        </script>
    </body>
    </html>
    """

from openpyxl.utils import get_column_letter

@app.route("/api/grade/stream")
def api_grade_stream():
    student_name = request.args.get('student_name')
    access_key = request.args.get('access_key')
    secret_key = request.args.get('secret_key')
    session_token = request.args.get('session_token')

    def generate():
        q = queue.Queue()
        def progress_callback(msg):
            q.put({"type": "progress", "message": msg})
            
        def run_validation():
            try:
                grader = InfrastructureGrader(access_key, secret_key, session_token, progress_callback=progress_callback)
                grader.run_all_checks()
                results = grader.results
                q.put({
                    "type": "complete",
                    "results": results,
                    "metadata": grader.discovery_metadata
                })
            except Exception as e:
                q.put({"type": "error", "message": str(e)})
            q.put(None) # Sentinel
            
        threading.Thread(target=run_validation).start()
        while True:
            item = q.get()
            if item is None: break
            yield f"data: {json.dumps(item)}\n\n"
            
    return Response(stream_with_context(generate()), mimetype='text/event-stream')

@app.route("/api/reset", methods=["POST"])
def api_reset():
    with open(DB_FILE, "w") as f:
        json.dump([], f)
    return jsonify({"status": "reset"})

@app.route("/api/delete/<id>", methods=["DELETE"])
def api_delete(id):
    db = get_db()
    db = [e for e in db if e['id'] != id]
    with open(DB_FILE, "w") as f:
        json.dump(db, f, indent=4)
    return jsonify({"status": "deleted"})

@app.route("/api/finalize", methods=["POST"])
def api_finalize():
    entry = request.json
    save_to_db(entry)
    return jsonify({"status": "finalized"})

@app.route("/api/leaderboard")
def api_leaderboard():
    db = get_db()
    # Sort by percentage descending
    sorted_db = sorted(db, key=lambda x: float(x['percentage']), reverse=True)
    
    # Calculate rank with ties (1, 1, 3...)
    results = []
    if not sorted_db: return jsonify([])
    
    current_rank = 1
    for i, entry in enumerate(sorted_db):
        if i > 0 and float(sorted_db[i]['percentage']) < float(sorted_db[i-1]['percentage']):
            current_rank = i + 1
        entry['rank'] = current_rank
        results.append(entry)
    
    return jsonify(results)

@app.route("/api/export", methods=["POST"])
def api_export():
    db = get_db()
    if not db: return jsonify({"status": "error", "message": "No data to export"}), 400
    
    # Sort students by percentage descending
    sorted_db = sorted(db, key=lambda x: float(x['percentage']), reverse=True)
    
    from openpyxl.comments import Comment

    # Style configuration
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    pass_font = Font(color="006100")
    pass_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    fail_font = Font(color="9C0006")
    fail_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    label_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    zebra_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    wb = Workbook()
    ws = wb.active
    ws.title = "Master Dashboard"
    
    # Detailed Log Sheet (All results)
    ws_log = wb.create_sheet("Detailed Evaluation Log")
    ws_log.append(["Student", "Category", "Metric", "Status", "Expected", "Actual", "Feedback"])
    for cell in ws_log[1]:
        cell.fill = header_fill
        cell.font = header_font

    # Identify all unique point NAMES (for the vertical table)
    all_item_names = []
    seen_names = set()
    for entry in sorted_db:
        for res in entry.get('results', []):
            full_name = f"{res['Category']} | {res['Item']}"
            if full_name not in seen_names:
                all_item_names.append(full_name)
                seen_names.add(full_name)
    
    # Sort all_item_names to ensure consistent category grouping
    all_item_names.sort(key=lambda x: x.split(" | ")[0])
    
    # Headers
    headers = ["Evaluation Point / Metric"] + [s['name'] for s in sorted_db]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    # Metadata Rows (DIRECTLY FROM DB)
    meta_rows = [
        ("Rank", lambda s, i: i + 1),
        ("Total Score", lambda s, i: s['total_score']),
        ("Max Possible", lambda s, i: s['max_score']),
        ("Percentage (%)", lambda s, i: float(s['percentage']))
    ]
    
    for label, func in meta_rows:
        row = [label]
        for idx, s in enumerate(sorted_db):
            row.append(func(s, idx))
        ws.append(row)
        curr_row = ws.max_row
        ws.cell(row=curr_row, column=1).fill = label_fill
        ws.cell(row=curr_row, column=1).font = Font(bold=True)
        if label == "Percentage (%)":
            for col in range(2, len(headers) + 1):
                ws.cell(row=curr_row, column=col).number_format = '0.00"%"'

    # Points Rows
    points_start_row = ws.max_row + 1
    current_cat = ""
    for idx_item, full_name in enumerate(all_item_names):
        cat_prefix = full_name.split(" | ")[0]
        if cat_prefix != current_cat:
            current_cat = cat_prefix
            ws.append([f"--- {current_cat} ---"] + ["" for _ in sorted_db])
            for cell in ws[ws.max_row]:
                cell.fill = label_fill
                cell.font = Font(bold=True, italic=True)

        row = [full_name]
        for s in sorted_db:
            # Find the result in the list (handle duplicates by taking the first match)
            res_obj = next((r for r in s.get('results', []) if f"{r['Category']} | {r['Item']}" == full_name), None)
            status = res_obj['Status'] if res_obj else "N/A"
            row.append(status)
        
        ws.append(row)
        curr_row = ws.max_row
        ws.cell(row=curr_row, column=1).font = Font(size=9)
        if idx_item % 2 == 0: ws.cell(row=curr_row, column=1).fill = zebra_fill
        
        for idx, s in enumerate(sorted_db):
            cell = ws.cell(row=curr_row, column=idx + 2)
            cell.alignment = Alignment(horizontal="center")
            if cell.value == "PASS":
                cell.font = pass_font
                cell.fill = pass_fill
            elif cell.value == "FAIL":
                cell.font = fail_font
                cell.fill = fail_fill
            else:
                if idx_item % 2 == 0: cell.fill = zebra_fill

    # Fill the Detailed Log Sheet
    for s in sorted_db:
        for res in s.get('results', []):
            ws_log.append([s['name'], res['Category'], res['Item'], res['Status'], res['Expected'], res['Actual'], res['Feedback']])

    # Formatting
    ws.column_dimensions['A'].width = 60
    ws.freeze_panes = "B6"
    for col in range(2, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 15

    ws_log.column_dimensions['A'].width = 20
    ws_log.column_dimensions['B'].width = 20
    ws_log.column_dimensions['C'].width = 40
    ws_log.column_dimensions['G'].width = 60

    filename = f"Master_Evaluation_Report_{datetime.now().strftime('%Y%m%d')}.xlsx"
    filepath = os.path.join(EXPORTS_DIR, filename)
    wb.save(filepath)
    return send_file(filepath, as_attachment=True)

    # Audit Sheet formatting
    ws_audit.column_dimensions['A'].width = 20
    ws_audit.column_dimensions['B'].width = 20
    ws_audit.column_dimensions['C'].width = 40
    ws_audit.column_dimensions['G'].width = 60

    # Main Sheet formatting
    ws.column_dimensions['A'].width = 60
    ws.freeze_panes = "B6" # Freeze header and first column
    for col in range(2, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 20

@app.route('/api/export-html', methods=['POST'])
def api_export_html():
    db = get_db()
    sorted_db = sorted(db, key=lambda x: float(x.get('percentage', 0)), reverse=True)
    
    html_template = f"""
<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>AWS Infrastructure Leaderboard</title>
    <link href="https://fonts.googleapis.com/css2?family=Plus+Jakarta+Sans:wght@300;400;600;700&display=swap" rel="stylesheet">
    <style>
        :root {{
            --primary: #6366f1; --secondary: #a855f7; --bg: #0f172a;
            --card: #1e293b; --text: #f8fafc; --success: #22c55e; --danger: #ef4444;
        }}
        * {{ margin: 0; padding: 0; box-sizing: border-box; font-family: 'Plus Jakarta Sans', sans-serif; }}
        body {{ background: var(--bg); color: var(--text); padding: 2rem; min-height: 100vh; }}
        .container {{ max-width: 1000px; margin: 0 auto; }}
        header {{ text-align: center; margin-bottom: 3rem; }}
        h1 {{ font-size: 2.5rem; background: linear-gradient(to right, var(--primary), var(--secondary)); -webkit-background-clip: text; -webkit-text-fill-color: transparent; }}
        .student-card {{ background: var(--card); border-radius: 1rem; padding: 1.5rem; display: flex; align-items: center; justify-content: space-between; margin-bottom: 1rem; cursor: pointer; transition: 0.3s; border: 1px solid rgba(255,255,255,0.05); }}
        .student-card:hover {{ transform: translateY(-5px); border-color: var(--primary); box-shadow: 0 10px 25px rgba(0,0,0,0.3); }}
        .rank {{ font-size: 1.5rem; font-weight: 700; width: 60px; color: rgba(255,255,255,0.2); }}
        .card-0 .rank {{ color: #ffd700; }} .card-1 .rank {{ color: #c0c0c0; }} .card-2 .rank {{ color: #cd7f32; }}
        .percentage {{ font-size: 1.5rem; font-weight: 700; color: var(--primary); text-align: right; }}
        #modal {{ display: none; position: fixed; top: 0; left: 0; width: 100%; height: 100%; background: rgba(15,23,42,0.95); backdrop-filter: blur(10px); z-index: 1000; padding: 2rem; overflow-y: auto; }}
        .modal-content {{ background: var(--card); max-width: 900px; margin: 0 auto; border-radius: 1.5rem; padding: 2rem; border: 1px solid rgba(255,255,255,0.1); }}
        .res-item {{ padding: 1rem; border-bottom: 1px solid rgba(255,255,255,0.05); display: flex; justify-content: space-between; }}
        .badge {{ padding: 4px 12px; border-radius: 20px; font-size: 0.75rem; font-weight: 700; }}
        .PASS {{ background: rgba(34,197,94,0.2); color: #4ade80; }}
        .FAIL {{ background: rgba(239,68,68,0.2); color: #f87171; }}
    </style>
</head>
<body>
    <div class="container">
        <header><h1>Infrastructure Leaderboard</h1><p>Cloud Computing Final Results</p></header>
        <div id="leaderboard"></div>
    </div>
    <div id="modal"><div class="modal-content"><div style="display:flex;justify-content:space-between;align-items:center;margin-bottom:2rem"><h2 id="mName"></h2><button onclick="closeM()" style="background:none;border:none;color:white;font-size:2rem;cursor:pointer">&times;</button></div><div id="mRes"></div></div></div>
    <script>
        const students = {json.dumps(sorted_db)};
        document.getElementById('leaderboard').innerHTML = students.map((s, i) => `
            <div class="student-card card-${{i}}" onclick="showM(${{i}})">
                <div class="rank">#${{i+1}}</div>
                <div style="flex-grow:1"><div style="font-size:1.2rem;font-weight:600">${{s.name}}</div><div style="font-size:0.8rem;opacity:0.5">${{s.total_score}} / ${{s.max_score}} Points</div></div>
                <div class="percentage">${{s.percentage}}%</div>
            </div>
        `).join('');
        function showM(i) {{
            const s = students[i]; document.getElementById('mName').innerText = s.name;
            document.getElementById('mRes').innerHTML = s.results.map(r => `
                <div class="res-item">
                    <div><div style="font-weight:600">${{r.Category}}</div><div style="font-size:0.8rem;opacity:0.7">${{r.Item}}</div><div style="font-size:0.75rem;opacity:0.5;margin-top:5px">${{r.Feedback}}</div></div>
                    <div><span class="badge ${{r.Status}}">${{r.Status}}</span></div>
                </div>
            `).join('');
            document.getElementById('modal').style.display = 'block';
        }}
        function closeM() {{ document.getElementById('modal').style.display = 'none'; }}
    </script>
</body>
</html>
    """
    filename = f"Leaderboard_Report_{datetime.now().strftime('%Y%m%d')}.html"
    filepath = os.path.join(EXPORTS_DIR, filename)
    with open(filepath, 'w', encoding='utf-8') as f:
        f.write(html_template)
    return send_file(filepath, as_attachment=True)

if __name__ == "__main__":
    app.run(debug=True, port=5000, threaded=True)
